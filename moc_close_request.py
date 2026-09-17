#!/usr/bin/env python3
"""BlackRock request: one row per MOC algo order, AU / HK / JP / SG, from the AWS extract.

Concatenates every AWS parquet, keeps the trade dates in the window, and
builds the CSV the request asks for:

    trade_date, sedol, bbg_ticker, ric, market, side, qty_order,
    qty_close_sent, qty_close_exec, px_cont_last, px_close, volume_close,
    unfilled_reason, snap_sent_bps, snap_used_bps, limit_status

plus qty_exec and qty_residual after qty_order, which the request did not
ask for.

snap_sent_bps, snap_used_bps and limit_status are NOT computed. They are
written as empty columns so the file keeps the requested layout.

Set AWS_DIR, KDB_HOST, KDB_PORT and CROSSCODE_PATH below first.

    python moc_close_request.py --probe          # check the mapping and STOP - run this first
    python moc_close_request.py                  # build the files, identifiers from CrossCode and kdb
    python moc_close_request.py --no-kdb         # build without kdb: no sedol
    python moc_close_request.py --no-crosscode   # build without CrossCode: SG ric may be blank
    python moc_close_request.py --self-test      # synthetic parquet and a fake kdb, end to end

Output, in --out (default output_moc_request/):

    aws_concat_<from>_<to>.csv    every parquet row in the window, one per aggrTgtId
    moc_close_<from>_<to>.csv     the file to send
    moc_close_<from>_<to>.xlsx    the same, plus the audit sheet
    moc_close_audit.csv           same orders with aggrTgtId, sym, the source
                                  columns and which rule set each value
    run_log.txt                   what was read, dropped, derived and assumed

What comes from where (AWS column -> requested column):

    _date            trade_date
    CrossCode        ric (RicCode), bbg_ticker (BloombergCode) where kdb has none
    kdb              sedol (ID_SEDOL1), bbg_ticker (TICKER + EQY_PRIM_EXCH_SHRT),
                     ric (ric_code) where CrossCode has none; equity_master,
                     else equity
    sym              ric / bbg_ticker where neither has one, converted where
                     the conversion is mechanical (JP, HK, AU)
    country / sym    market
    side             side (SSH -> SELL_SHORT)
    ordqty           qty_order
    cumqty           qty_exec, executed over the whole order
    derived          qty_residual = qty_order - qty_exec, never below 0
    derived          qty_close_sent, see CLOSE_SENT_SOURCE
    fillCloseSize    qty_close_exec
    last_cont_price  px_cont_last
    endprice         px_close  (PX_LAST is not the day's close - see moc_tca.py)
    marketCloseSize  volume_close
    derived          unfilled_reason, see assign_unfilled_reason
"""
from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd

import moc_tca as tca

# --- where the data is ----------------------------------------------------
# Folder holding the AWS parquet files. Searched recursively.
AWS_DIR = r"data\aws"

# The kdb server holding equity_master and equity - the same host and port
# the kdb-queries scripts use for those two tables.
KDB_HOST = "CHANGEME"
KDB_PORT = 0

# The security master, the same CrossCode.csv Nova's LimitUpDown reads.
# ric comes from its RicCode.
CROSSCODE_PATH = r"CHANGEME\CrossCode.csv"

# --- scope ----------------------------------------------------------------
DATE_FROM = "2026-07-11"
DATE_TO = "2026-09-11"
MARKETS = ["AU", "HK", "JP", "SP"]

# Which orders are "MOC algo orders". Matched on the AWS algo column,
# upper-cased. On this platform the CLOSE label covers the MOC product
# (moc_tca.py scopes the same way). --probe prints every value with its
# count and auction share: check it before sending anything.
ALGO_COLUMNS = ["clientAlgo", "tag9001"]
MOC_ALGOS = {"CLOSE", "MOC"}

# --- qty_close_sent -------------------------------------------------------
#   "residual"  ordqty minus everything executed outside the closing auction,
#               i.e. what was still working when the auction started. Zero if
#               the order ended before the continuous session did.
#   "qatt"      maxQattCloseSym as it stands. Use it once the data owner
#               confirms it means quantity sent to the close. --probe prints
#               the two side by side.
CLOSE_SENT_SOURCE = "residual"

# --- unfilled_reason ------------------------------------------------------
# Short-sale price rules that can stop a short sell printing in the auction.
# AU and SG have no uptick rule, so they are not listed.
#   JP  price restriction is on once the stock trades 10% under the previous
#       close; a short sell then cannot print at or below the last price.
#   HK  a short sell in the closing auction cannot be below the reference
#       price. last_cont_price stands in for the reference price.
# Both are approximations from end-of-day fields. Confirm with the desk.
JP_RESTRICTION_DROP = 0.10

# Set to the auction participation cap in percent (e.g. 20.0) if the algo
# applies one. An unfilled order whose auction fill reached the cap, within
# VOLUME_CAP_TOL, is then coded VOLUME_CAP. None = the cap is unknown and the
# code is never assigned.
AUCTION_VOLUME_CAP_PCT: float | None = None
VOLUME_CAP_TOL = 0.05

# Written for unfilled orders no rule explains. Every one is listed in the
# audit file; replace them before the file goes out.
OTHER_UNEXPLAINED = "OTHER: cause not identified from execution data"

OUTPUT_COLUMNS = ["trade_date", "sedol", "bbg_ticker", "ric", "market", "side",
                  "qty_order", "qty_exec", "qty_residual",
                  "qty_close_sent", "qty_close_exec",
                  "px_cont_last", "px_close", "volume_close", "unfilled_reason",
                  "snap_sent_bps", "snap_used_bps", "limit_status"]
NOT_COMPUTED = ["snap_sent_bps", "snap_used_bps", "limit_status"]

# AWS names, first hit wins. Matched exactly, then case-insensitively.
# Not through moc_tca._norm_name: that strips underscores, so "_date" and
# "date", "side" and "side_", "country" and "_country" would all collide.
SOURCE = {
    "id":           ["aggrTgtId"],
    "date":         ["_date", "date", "Date"],
    "sym":          ["sym", "Sym"],
    "country":      ["country", "_country"],
    "side":         ["side", "side_"],
    "ordqty":       ["ordqty"],
    "cumqty":       ["cumqty"],
    "fill_close":   ["fillCloseSize"],
    "market_close": ["marketCloseSize"],
    "px_cont_last": ["last_cont_price"],
    "px_close":     ["endprice"],
    "limit":        ["ordprice"],
    "prev_close":   ["PreviousClose"],
    "px_low":       ["PX_LOW"],
    "end_time":     ["fend_time", "end_time"],
    "qatt_close":   ["maxQattCloseSym"],
    "sedol":        ["sedol", "SEDOL"],
    "ric":          ["ric", "RIC"],
    "bbg":          ["bbg_ticker", "bbgTicker", "BBG"],
}
NEEDED = ["id", "date", "sym", "side", "ordqty", "cumqty", "fill_close",
          "market_close", "px_cont_last", "px_close"]

SIDE_MAP = {
    "B": "BUY", "BUY": "BUY", "BOT": "BUY", "1": "BUY",
    "S": "SELL", "SELL": "SELL", "SLD": "SELL", "2": "SELL",
    "SS": "SELL_SHORT", "SSH": "SELL_SHORT", "SSE": "SELL_SHORT",
    "SHORT": "SELL_SHORT", "SHORTSELL": "SELL_SHORT", "SELLSHORT": "SELL_SHORT",
    "5": "SELL_SHORT",
    "BC": "BUY_COVER", "BTC": "BUY_COVER", "COVER": "BUY_COVER",
    "BUYCOVER": "BUY_COVER", "BUYTOCOVER": "BUY_COVER",
}

COUNTRY_MAP = {
    "AU": "AU", "AUS": "AU", "AUSTRALIA": "AU",
    "HK": "HK", "HKG": "HK", "HONGKONG": "HK",
    "JP": "JP", "JPN": "JP", "JAPAN": "JP",
    "SG": "SG", "SGP": "SG", "SINGAPORE": "SG",
}
RIC_SUFFIX = {"AX": "AU", "T": "JP", "SI": "SG"}
# The house sym is a Bloomberg ticker dot-joined to its exchange code
# (BHP.AU, 9984.JP), so these are read after either a dot or a space.
# HK is the same code on both sides and is handled by the code itself.
BBG_SUFFIX = {"AU": "AU", "AT": "AU", "HK": "HK", "JP": "JP", "JT": "JP",
              "SP": "SG"}
BBG_EXCHANGE = {"AU": "AT", "HK": "HK", "JP": "JT", "SG": "SP"}
BBG_COMPOSITE = {"AU": "AU", "HK": "HK", "JP": "JP", "SG": "SP"}
RIC_EXCHANGE = {"AU": "AX", "HK": "HK", "JP": "T", "SG": "SI"}
MARKET_NAME = {"AU": "Australia", "HK": "Hong Kong", "JP": "Japan",
               "SG": "Singapore"}

_LOG: list[str] = []


def log(msg: str = "") -> None:
    print(msg)
    _LOG.append(str(msg))


def warn(msg: str) -> None:
    log(f"  !! {msg}")


# ===========================================================================
# READING
# ===========================================================================

def find_column(columns, candidates) -> str | None:
    cols = list(columns)
    for cand in candidates:
        if cand in cols:
            return cand
    lower = {str(c).lower(): c for c in cols}
    for cand in candidates:
        if cand.lower() in lower:
            return lower[cand.lower()]
    return None


def parse_dates(s: pd.Series) -> pd.Series:
    """_date as datetime, date, 20260721, '2026.07.21' or '2026-07-21'."""
    if pd.api.types.is_datetime64_any_dtype(s):
        return s.dt.normalize()
    if pd.api.types.is_numeric_dtype(s):
        return pd.to_datetime(s.astype("Int64").astype(str), format="%Y%m%d",
                              errors="coerce")
    txt = s.astype(str).str.strip().str.replace(".", "-", regex=False)
    return pd.to_datetime(txt, format="mixed", errors="coerce").dt.normalize()


def date_from_filename(path: Path):
    m = re.search(r"(20\d{2})[-._]?(\d{2})[-._]?(\d{2})", path.name)
    return pd.Timestamp(f"{m[1]}-{m[2]}-{m[3]}") if m else pd.NaT


def read_window(folder: Path, d_from: pd.Timestamp, d_to: pd.Timestamp) -> pd.DataFrame:
    """Every parquet under folder, rows in [d_from, d_to] only, concatenated."""
    files = sorted(folder.rglob("*.parquet"))
    log(f"  {len(files)} parquet file(s) under {folder}")
    frames, unreadable, no_date = [], 0, 0
    rows_read = 0
    for path in files:
        try:
            df = pd.read_parquet(path)
        except Exception as exc:
            warn(f"could not read {path.name}: {exc}")
            unreadable += 1
            continue
        rows_read += len(df)
        col = find_column(df.columns, SOURCE["date"])
        if col is not None:
            dates = parse_dates(df[col])
        else:
            no_date += 1
            dates = pd.Series(date_from_filename(path), index=df.index)
        keep = dates.between(d_from, d_to)
        if keep.any():
            df = df.loc[keep].copy()
            df["trade_date_parsed"] = dates[keep]
            df["source_file"] = path.name
            frames.append(df)
    if unreadable:
        warn(f"{unreadable} file(s) could not be read and are left out")
    if no_date:
        warn(f"{no_date} file(s) had no date column; the date was taken from "
             f"the file name")
    log(f"  {rows_read:,} rows read, "
        f"{sum(len(f) for f in frames):,} in {d_from:%Y-%m-%d} .. {d_to:%Y-%m-%d}")
    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True, sort=False)

    key = find_column(out.columns, SOURCE["id"])
    if key is not None:
        dupes = int(out[key].duplicated().sum())
        if dupes:
            warn(f"{dupes:,} duplicate {key} values - keeping the first of each")
            out = out.drop_duplicates(subset=[key], keep="first")
    return out.reset_index(drop=True)


def resolve(frame: pd.DataFrame) -> dict[str, str]:
    return {k: c for k, cands in SOURCE.items()
            if (c := find_column(frame.columns, cands)) is not None}


# ===========================================================================
# DERIVATIONS
# ===========================================================================

def market_of(frame: pd.DataFrame, cols: dict) -> tuple[pd.Series, int]:
    """AU/HK/JP/SG from country, falling back to the symbol suffix."""
    from_sym = frame[cols["sym"]].map(sym_market)
    if "country" not in cols:
        return from_sym, 0
    key = (frame[cols["country"]].astype(str).str.upper()
           .str.replace(r"[\s_.-]", "", regex=True))
    from_country = key.map(COUNTRY_MAP)
    disagree = int((from_country.notna() & from_sym.notna()
                    & (from_country != from_sym)).sum())
    return from_country.fillna(from_sym), disagree


def sym_parts(sym) -> tuple[str | None, str | None, str | None]:
    """(code, market, style) from '9984.T' (ric), '9984.JP' or '9984 JT' (bbg)."""
    if not isinstance(sym, str) or not sym.strip():
        return None, None, None
    s = sym.strip().upper()
    m = re.fullmatch(r"(\S+?)\.([A-Z]{1,2})", s)
    if m and m[2] in RIC_SUFFIX:
        return m[1], RIC_SUFFIX[m[2]], "ric"
    m = re.fullmatch(r"(\S+?)[. ]\s*([A-Z]{2})(?:\s+EQUITY)?", s)
    if m and m[2] in BBG_SUFFIX:
        return m[1], BBG_SUFFIX[m[2]], "bbg"
    return None, None, None


def sym_market(sym):
    return sym_parts(sym)[1]


def scope_markets() -> set[str]:
    """MARKETS as written, country or Bloomberg code: SP and SG are both Singapore."""
    return {COUNTRY_MAP.get(m.upper(), BBG_SUFFIX.get(m.upper(), m.upper()))
            for m in MARKETS}


def identifiers(sym) -> tuple[str | None, str | None]:
    """(ric, bbg_ticker) from the symbol, where the conversion is mechanical.

    JP and AU codes are the same on both. HK drops leading zeros on
    Bloomberg and pads to four digits on Reuters. SG stock codes (D05) are
    not Bloomberg tickers (DBS), so only the side already held is kept.
    """
    code, mkt, style = sym_parts(sym)
    if code is None:
        return None, None
    if mkt == "HK":
        if not code.isdigit():
            return None, None
        return f"{int(code):04d}.HK", f"{int(code)} HK"
    ric = bbg = None
    if style == "ric":
        ric = f"{code}.{RIC_EXCHANGE[mkt]}"
        if mkt in ("JP", "AU"):
            bbg = f"{code} {BBG_EXCHANGE[mkt]}"
    else:
        bbg = f"{code} {BBG_EXCHANGE[mkt]}"
        if mkt in ("JP", "AU"):
            ric = f"{code}.{RIC_EXCHANGE[mkt]}"
    return ric, bbg


def num(frame: pd.DataFrame, cols: dict, key: str) -> pd.Series:
    if key not in cols:
        return pd.Series(np.nan, index=frame.index)
    return tca._to_num(frame[cols[key]])


def ended_before_close(frame, cols, market, dates) -> pd.Series:
    """True where the order ended before its market's continuous session did.

    fend_time is read as HKT, like fstart_time in moc_tca.py. If that is wrong
    the test would fire on filled orders too, so it checks itself on orders
    that did print in the auction and switches off if they fail it.
    """
    none = pd.Series(False, index=frame.index)
    if "end_time" not in cols:
        warn("no fend_time / end_time: an order cancelled before the close "
             "cannot be told apart, so its residual counts as sent")
        return none
    end_min = tca._parse_time(frame[cols["end_time"]])
    close_min = tca.continuous_end_min(market.map(MARKET_NAME), dates)
    early = (end_min < close_min).fillna(False)
    printed = num(frame, cols, "fill_close").fillna(0) > 0
    if printed.sum() >= 20 and (early & printed).sum() > 0.2 * printed.sum():
        warn(f"{int((early & printed).sum()):,} of {int(printed.sum()):,} orders "
             f"that printed in the auction appear to end before the close. The "
             f"end time is not HKT, so this test is switched off.")
        return none
    return early


def close_sent(frame, cols, ordqty, cumqty, fill, early) -> pd.Series:
    if CLOSE_SENT_SOURCE == "qatt":
        if "qatt_close" not in cols:
            raise SystemExit("CLOSE_SENT_SOURCE = 'qatt' but there is no "
                             "maxQattCloseSym column")
        return num(frame, cols, "qatt_close")
    residual = (ordqty - (cumqty - fill)).clip(lower=0)
    residual = np.maximum(residual, fill)
    residual = residual.clip(upper=ordqty.where(ordqty.notna(), np.inf))
    return residual.where(~(early & (fill <= 0)), 0.0)


def assign_unfilled_reason(d: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    """(unfilled_reason, rule) per order. Blank reason = fully filled.

    d holds side, market, qty_order, cumqty, qty_close_sent, qty_close_exec,
    px_cont_last, px_close, volume_close, limit, prev_close, px_low, early.
    First rule that fires wins, in this order:

      data missing          quantities absent - left blank, flagged in audit
      filled                nothing left unfilled in the auction
      ended before close    nothing sent, order ended before the auction
      no auction            volume_close is 0
      LIMIT_QUEUE           auction printed exactly at the order's limit
      limit did not cross   buy limit under / sell limit over the auction price
      VOLUME_CAP            auction fill reached AUCTION_VOLUME_CAP_PCT
      UPTICK_RULE           short sell caught by the JP or HK price rule
      unexplained           none of the above
    """
    reason = pd.Series("", index=d.index, dtype=object)
    rule = pd.Series("", index=d.index, dtype=object)
    open_ = pd.Series(True, index=d.index)

    def put(mask, code, name):
        nonlocal open_
        hit = mask.fillna(False) & open_
        reason[hit] = code
        rule[hit] = name
        open_ = open_ & ~hit

    sent, execd = d["qty_close_sent"], d["qty_close_exec"]
    put(sent.isna() | execd.isna() | d["qty_order"].isna(), "", "data missing")
    put((sent <= 0) & (d["cumqty"] >= d["qty_order"]), "", "filled")
    put((sent > 0) & (execd >= sent), "", "filled")
    put((sent <= 0) & d["early"],
        "OTHER: order ended before the closing auction", "ended before close")
    put(sent <= 0, "OTHER: nothing sent to the closing auction", "nothing sent")
    put(d["volume_close"] <= 0, "OTHER: no closing auction held", "no auction")

    lim, px = d["limit"], d["px_close"]
    has_lim = lim > 0
    at_limit = pd.Series(np.isclose(lim, px, rtol=1e-7, atol=1e-9), index=d.index)
    put(has_lim & at_limit, "LIMIT_QUEUE", "LIMIT_QUEUE")
    buy = d["side"].isin(["BUY", "BUY_COVER"])
    put(has_lim & ((buy & (lim < px)) | (~buy & (lim > px))),
        "OTHER: limit price did not cross the auction price", "limit did not cross")

    if AUCTION_VOLUME_CAP_PCT is not None:
        share = 100.0 * execd / d["volume_close"]
        put(share >= AUCTION_VOLUME_CAP_PCT * (1 - VOLUME_CAP_TOL),
            "VOLUME_CAP", "VOLUME_CAP")

    short = d["side"].eq("SELL_SHORT")
    jp_on = d["px_low"] <= (1 - JP_RESTRICTION_DROP) * d["prev_close"]
    put(short & d["market"].eq("JP") & jp_on & (px <= d["px_cont_last"]),
        "UPTICK_RULE", "UPTICK_RULE (JP price restriction)")
    put(short & d["market"].eq("HK") & (px < d["px_cont_last"]),
        "UPTICK_RULE", "UPTICK_RULE (HK short sell below reference price)")

    put(pd.Series(True, index=d.index), OTHER_UNEXPLAINED, "unexplained")
    return reason, rule


# ===========================================================================
# KDB - identifiers from equity_master / equity
# ===========================================================================
# Tried in order; the first that answers with rows is used. equity carries no
# ric_code, so off that table ric stays as derived from the sym.
KDB_TABLES = {
    "equity_master": ["ID_SEDOL1", "TICKER", "EQY_PRIM_EXCH_SHRT", "ric_code"],
    "equity":        ["ID_SEDOL1", "TICKER", "EQY_PRIM_EXCH_SHRT"],
}
KDB_CHUNK = 500
Q_EPOCH = pd.Timestamp("2000-01-01")


def kdb_query(table: str, fields: list[str]) -> str:
    # Dates go over the wire as day counts and are cast in q: a python date
    # can arrive as the wrong q type and die on 'type (Nova, 2026-09-04).
    # No `$ on s and no `by`: pykx sends symbols already, and a keyed
    # result needs a q licence to read back.
    return ("{[a;b;s] select date,sym," + ",".join(fields) + " from " + table
            + " where date within (\"d\"$a;\"d\"$b), sym in s}")


def kdb_connect():
    if KDB_HOST == "CHANGEME" or not KDB_PORT:
        raise SystemExit("ERROR: set KDB_HOST and KDB_PORT at the top of "
                         "moc_close_request.py, or run with --no-kdb")
    try:
        import pykx
    except ImportError:
        raise SystemExit("ERROR: pykx is not installed (pip install pykx), "
                         "or run with --no-kdb")
    log(f"  connecting to kdb {KDB_HOST}:{KDB_PORT}")
    return pykx.SyncQConnection(host=KDB_HOST, port=int(KDB_PORT))


def kdb_candidates(sym, extra=()) -> list[str]:
    """The sym as the extract spells it, then the kdb form: TICKER.COMPOSITE.

    extra is the CrossCode match for the row, which is the only way to reach
    a name whose sym does not convert (an SG RIC to its Bloomberg ticker).
    """
    s = str(sym).strip()
    out = [s]
    code, mkt, _ = sym_parts(s)
    if code is not None:
        alts = [f"{code}.{BBG_COMPOSITE[mkt]}"]
        if mkt == "HK" and code.isdigit():
            alts = [f"{int(code)}.HK", f"{int(code):04d}.HK"]
        out += [a for a in alts if a not in out]
    for e in extra:
        for a in kdb_candidates(e):
            if a and a not in out:
                out.append(a)
    return out


def row_candidates(audit: pd.DataFrame) -> list[list[str]]:
    extras = [c for c in ("cc_fidessa", "cc_bbg") if c in audit]
    out = []
    for i in audit.index:
        extra = [audit.at[i, c] for c in extras
                 if isinstance(audit.at[i, c], str) and audit.at[i, c]]
        out.append(kdb_candidates(audit.at[i, "sym"], extra))
    return out


def _kdb_frame(result) -> pd.DataFrame:
    df = result.pd() if hasattr(result, "pd") else pd.DataFrame(result)
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].map(lambda v: v.decode() if isinstance(v, (bytes, bytearray)) else v)
    return df


def fetch_equity(conn, d_from, d_to, syms: list[str]) -> tuple[pd.DataFrame, str]:
    a, b = int((d_from - Q_EPOCH).days), int((d_to - Q_EPOCH).days)
    for table, fields in KDB_TABLES.items():
        query = kdb_query(table, fields)
        try:
            frames = [_kdb_frame(conn(query, a, b, syms[i:i + KDB_CHUNK]))
                      for i in range(0, len(syms), KDB_CHUNK)]
        except Exception as exc:                            # noqa: BLE001
            warn(f"{table}: {type(exc).__name__}: {exc}")
            log(f"    query: {query}")
            continue
        eq = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
        found = eq["sym"].nunique() if len(eq) else 0
        log(f"  {table}: {len(eq):,} rows, {found:,} of {len(syms):,} syms asked")
        if len(eq):
            return eq, table
    raise SystemExit(
        "ERROR: no identifiers came back from kdb. If a table answered with 0 "
        "rows, the sym form is wrong rather than the data missing - asked for "
        f"e.g. {syms[:4]}")


def _set_identifier(out, audit, col, value: pd.Series, source: str,
                    keep: pd.Series | None = None) -> None:
    """Take value where it has one, unless keep says the current one stays."""
    value = value.where(value.notna() & value.astype(str).str.strip().ne(""))
    use = value.notna() if keep is None else value.notna() & ~keep
    differ = use & out[col].notna() & (out[col] != value)
    if differ.any():
        ex = pd.DataFrame({"sym": audit.loc[differ, "sym"], "was": out.loc[differ, col],
                           source: value[differ]}).head(5)
        log(f"  {col}: {source} replaces a different value on "
            f"{int(differ.sum()):,} orders, e.g.\n{ex.to_string()}")
    out[col] = out[col].astype(object)
    audit[col] = audit[col].astype(object)
    out.loc[use, col] = value[use]
    audit.loc[use, col] = value[use]
    audit.loc[use, f"{col}_source"] = source


def apply_kdb(out: pd.DataFrame, audit: pd.DataFrame, eq: pd.DataFrame,
              table: str) -> None:
    """sedol and bbg_ticker from kdb; ric only where CrossCode had none.

    Row for the trade date first; failing that the latest row for the sym in
    the window, flagged in the audit as another date.
    """
    def text(col):
        if col not in eq:
            return pd.Series(None, index=eq.index, dtype=object)
        s = eq[col].astype(object).where(eq[col].notna())
        s = s.map(lambda v: str(v).strip() if v is not None else None)
        return s.where(s.ne("") & s.ne("nan") & s.notna())

    eq = eq.assign(kdb_date=pd.to_datetime(eq["date"]).dt.normalize(),
                   sym=eq["sym"].astype(str).str.strip())
    tick, exch = text("TICKER"), text("EQY_PRIM_EXCH_SHRT")
    eq = eq.assign(k_sedol=text("ID_SEDOL1"), k_ric=text("ric_code"),
                   k_bbg=(tick + " " + exch).where(tick.notna() & exch.notna()))
    eq = eq.sort_values("kdb_date")
    exact = eq.drop_duplicates(["kdb_date", "sym"], keep="last").set_index(["kdb_date", "sym"])
    latest = eq.drop_duplicates("sym", keep="last").set_index("sym")

    got = {k: [] for k in ("kdb_sym", "kdb_match", "k_sedol", "k_bbg", "k_ric")}
    days = pd.to_datetime(out["trade_date"])
    for cands, day in zip(row_candidates(audit), days):
        hit, sym, how = None, None, "not found"
        for cand in cands:
            if (day, cand) in exact.index:
                hit, sym, how = exact.loc[(day, cand)], cand, "trade date"
                break
        if hit is None:
            for cand in cands:
                if cand in latest.index:
                    hit, sym, how = latest.loc[cand], cand, "other date"
                    break
        got["kdb_sym"].append(sym)
        got["kdb_match"].append(how)
        for k in ("k_sedol", "k_bbg", "k_ric"):
            got[k].append(None if hit is None else hit[k])

    k = pd.DataFrame(got, index=out.index)
    matched = k["kdb_match"].value_counts()
    log(f"  {table} matched: " + ", ".join(f"{n} {v:,}" for n, v in matched.items()))
    as_is = int((k["kdb_sym"] == audit["sym"].astype(str).str.strip()).sum())
    log(f"  kdb sym spelled as in the extract on {as_is:,}, converted on "
        f"{int(k['kdb_sym'].notna().sum()) - as_is:,}")

    from_cc = (audit["ric_source"].eq("crosscode") if "ric_source" in audit
               else pd.Series(False, index=out.index))
    _set_identifier(out, audit, "sedol", k["k_sedol"], "kdb")
    _set_identifier(out, audit, "bbg_ticker", k["k_bbg"], "kdb")
    _set_identifier(out, audit, "ric", k["k_ric"], "kdb", keep=from_cc)
    audit["kdb_sym"] = k["kdb_sym"]
    audit["kdb_match"] = k["kdb_match"]
    missing = k["kdb_match"].eq("not found")
    if missing.any():
        warn(f"{int(missing.sum()):,} orders not in {table}; their syms: "
             f"{sorted(audit.loc[missing, 'sym'].astype(str).unique())[:10]}")


# ===========================================================================
# CROSSCODE - the security master
# ===========================================================================
# Primary listings only: a Japanese name also has lines on JNX-MAIN and
# CHJ-MAIN, whose RICs are not the one the request wants.
CC_VENUES = {"ASX-MAIN": "AU", "HKG-MAIN": "HK", "HKG-GEM": "HK",
             "TYO-MAIN": "JP", "SES-MAIN": "SG"}
# LimitUpDown reads "FidessaCode", TradingData's copy has "#FidessaCode".
CC_FIDESSA = ("FidessaCode", "#FidessaCode")


def load_crosscode(path) -> pd.DataFrame:
    p = Path(path)
    if "CHANGEME" in str(path) or not p.is_file():
        raise SystemExit(f"ERROR: CROSSCODE_PATH is not a file: {path}. Set it at "
                         f"the top of moc_close_request.py, or run with --no-crosscode")
    cc = pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    cc.columns = [c.strip() for c in cc.columns]
    missing = [c for c in ("RicCode", "BloombergCode") if c not in cc.columns]
    if missing:
        raise SystemExit(f"ERROR: {p.name} has no {', '.join(missing)} column")
    fid = next((c for c in CC_FIDESSA if c in cc.columns), None)
    cc = cc.apply(lambda s: s.str.strip())
    cc["cc_fidessa"] = cc[fid] if fid else ""
    total = len(cc)

    bbg = cc["BloombergCode"].str.upper()
    cc["ticker"] = bbg.str.rsplit(" ", n=1).str[0]
    if "FidessaMarket" in cc:
        cc["market"] = cc["FidessaMarket"].map(CC_VENUES)
    else:
        warn(f"{p.name} has no FidessaMarket: markets taken from the "
             f"BloombergCode, so secondary venues cannot be told apart")
        cc["market"] = bbg.str.rsplit(" ", n=1).str[-1].map(BBG_SUFFIX)
    cc = cc[cc["market"].isin(scope_markets())
            & (cc["RicCode"].ne("") | cc["BloombergCode"].ne(""))].copy()

    # Several lines for one code: the live one wins, then the file's order.
    if "BloombergStatus" in cc:
        status = cc["BloombergStatus"].str.upper()
        cc["_rank"] = np.where(status.isin(["", "ACTV"]), 0, 1)
        cc = cc.sort_values("_rank", kind="stable")
    log(f"  CrossCode {p.name}: {total:,} rows, {len(cc):,} primary listings in "
        f"{'/'.join(sorted(scope_markets()))}")
    return cc


def apply_crosscode(out: pd.DataFrame, audit: pd.DataFrame, cc: pd.DataFrame) -> None:
    """ric (RicCode) and bbg_ticker (BloombergCode) from the CrossCode line.

    The sym is tried as a FidessaCode, then as a RicCode, then its ticker
    within its market against the BloombergCode ticker. The audit says which.
    """
    def index(key: pd.Series):
        k = key.str.upper()
        return cc[k.ne("")].assign(_k=k[k.ne("")]).drop_duplicates("_k").set_index("_k")

    by_fid = index(cc["cc_fidessa"])
    by_ric = index(cc["RicCode"])
    by_tick = index(cc["ticker"] + "|" + cc["market"])

    rows = {k: [] for k in ("cc_match", "cc_fidessa", "RicCode", "BloombergCode")}
    for sym, mkt, ric in zip(audit["sym"], out["market"], out["ric"]):
        s = str(sym).strip().upper()
        code = sym_parts(s)[0]
        if code is not None and mkt == "HK" and code.isdigit():
            code = str(int(code))
        tries = [("FidessaCode", by_fid, s), ("RicCode", by_ric, s)]
        if isinstance(ric, str):
            tries.append(("RicCode from sym", by_ric, ric.upper()))
        if code is not None:
            tries.append(("ticker", by_tick, f"{code}|{mkt}"))
        hit, how = None, "not found"
        for name, idx, key in tries:
            if key in idx.index and idx.loc[key, "market"] == mkt:
                hit, how = idx.loc[key], name
                break
        rows["cc_match"].append(how)
        for c in ("cc_fidessa", "RicCode", "BloombergCode"):
            rows[c].append(None if hit is None else hit[c])

    m = pd.DataFrame(rows, index=out.index)
    log("  CrossCode matched by: " + ", ".join(
        f"{n} {v:,}" for n, v in m["cc_match"].value_counts().items()))
    audit["cc_match"] = m["cc_match"]
    audit["cc_fidessa"] = m["cc_fidessa"]
    audit["cc_bbg"] = m["BloombergCode"]
    _set_identifier(out, audit, "ric", m["RicCode"], "crosscode")
    _set_identifier(out, audit, "bbg_ticker", m["BloombergCode"], "crosscode")
    missing = m["cc_match"].eq("not found")
    if missing.any():
        warn(f"{int(missing.sum()):,} orders not in the CrossCode; their syms: "
             f"{sorted(audit.loc[missing, 'sym'].astype(str).unique())[:10]}")


# ===========================================================================
# BUILD
# ===========================================================================

def build(aws: pd.DataFrame, cols: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    missing = [k for k in NEEDED if k not in cols]
    if missing:
        names = ", ".join(SOURCE[k][0] for k in missing)
        raise SystemExit(f"ERROR: the AWS extract has no {names}")

    frame = aws
    market, disagree = market_of(frame, cols)
    if disagree:
        warn(f"{disagree:,} orders have a country that disagrees with the "
             f"symbol suffix; country was used")

    algo_col = find_column(frame.columns, ALGO_COLUMNS)
    if algo_col is None:
        raise SystemExit(f"ERROR: no algo column ({', '.join(ALGO_COLUMNS)})")
    algo = frame[algo_col].astype(str).str.strip().str.upper()
    is_moc = algo.isin(MOC_ALGOS)
    in_mkt = market.isin(scope_markets())
    log(f"  {algo_col} in {sorted(MOC_ALGOS)}: {int(is_moc.sum()):,} orders; "
        f"in {'/'.join(MARKETS)}: {int(in_mkt.sum()):,}; both: "
        f"{int((is_moc & in_mkt).sum()):,}")
    if not (is_moc & in_mkt).any():
        raise SystemExit("ERROR: no MOC orders in scope. Run --probe and check "
                         "MOC_ALGOS against the algo values.")
    frame = frame.loc[is_moc & in_mkt].reset_index(drop=True)
    market = market[is_moc & in_mkt].reset_index(drop=True)

    raw_side = (frame[cols["side"]].astype(str).str.upper()
                .str.replace(r"[\s_-]", "", regex=True))
    side = raw_side.map(SIDE_MAP)
    unknown = raw_side[side.isna()].value_counts()
    if len(unknown):
        raise SystemExit("ERROR: side codes with no mapping - add them to "
                         "SIDE_MAP:\n" + unknown.to_string())

    dates = frame["trade_date_parsed"]
    ordqty = num(frame, cols, "ordqty")
    cumqty = num(frame, cols, "cumqty")
    volume_close = num(frame, cols, "market_close")
    fill = num(frame, cols, "fill_close")
    blank_fill = fill.isna() & (volume_close > 0)
    if blank_fill.any():
        log(f"  {int(blank_fill.sum()):,} orders have an empty fillCloseSize on a "
            f"day the auction ran; read as 0")
    fill = fill.where(~blank_fill, 0.0)

    early = ended_before_close(frame, cols, market, dates)
    sent = close_sent(frame, cols, ordqty, cumqty, fill, early)

    ids = frame[cols["sym"]].map(identifiers)
    ric = ids.map(lambda t: t[0])
    bbg = ids.map(lambda t: t[1])
    if "ric" in cols:
        ric = frame[cols["ric"]].where(frame[cols["ric"]].notna(), ric)
    if "bbg" in cols:
        bbg = frame[cols["bbg"]].where(frame[cols["bbg"]].notna(), bbg)
    sedol = frame[cols["sedol"]] if "sedol" in cols else pd.Series(None, index=frame.index, dtype=object)

    work = pd.DataFrame({
        "side": side, "market": market, "qty_order": ordqty, "cumqty": cumqty,
        "qty_close_sent": sent, "qty_close_exec": fill,
        "px_cont_last": num(frame, cols, "px_cont_last"),
        "px_close": num(frame, cols, "px_close"),
        "volume_close": volume_close, "limit": num(frame, cols, "limit"),
        "prev_close": num(frame, cols, "prev_close"),
        "px_low": num(frame, cols, "px_low"), "early": early,
    })
    reason, rule = assign_unfilled_reason(work)

    out = pd.DataFrame({
        "trade_date": dates.dt.strftime("%Y-%m-%d"),
        "sedol": sedol, "bbg_ticker": bbg, "ric": ric, "market": market,
        "side": side,
        "qty_order": ordqty.round().astype("Int64"),
        "qty_exec": cumqty.round().astype("Int64"),
        "qty_residual": (ordqty - cumqty).clip(lower=0).round().astype("Int64"),
        "qty_close_sent": sent.round().astype("Int64"),
        "qty_close_exec": fill.round().astype("Int64"),
        "px_cont_last": work["px_cont_last"], "px_close": work["px_close"],
        "volume_close": volume_close.round().astype("Int64"),
        "unfilled_reason": reason,
    })
    for c in NOT_COMPUTED:
        out[c] = ""
    out = out[OUTPUT_COLUMNS]

    audit = out.copy()
    audit.insert(0, "aggrTgtId", frame[cols["id"]])
    audit.insert(1, "sym", frame[cols["sym"]])
    audit.insert(2, algo_col, frame[algo_col])
    audit["side_raw"] = frame[cols["side"]]
    audit["ordprice"] = work["limit"]
    audit["PreviousClose"] = work["prev_close"]
    audit["PX_LOW"] = work["px_low"]
    if "end_time" in cols:
        audit["end_time"] = frame[cols["end_time"]]
    if "qatt_close" in cols:
        audit["maxQattCloseSym"] = num(frame, cols, "qatt_close")
    audit["ended_before_close"] = early
    for col in ("sedol", "bbg_ticker", "ric"):
        audit[f"{col}_source"] = np.where(out[col].notna(), "extract", "")
    audit["reason_rule"] = rule
    audit["source_file"] = frame["source_file"]

    order = ["trade_date", "market", "ric", "bbg_ticker"]
    keep = out.sort_values(order, kind="stable").index
    return out.loc[keep].reset_index(drop=True), audit.loc[keep].reset_index(drop=True)


def report(out: pd.DataFrame, audit: pd.DataFrame) -> None:
    tca_section("RESULT")
    log(f"  {len(out):,} orders")
    log("")
    g = out.groupby("market")
    log(pd.DataFrame({
        "orders": g.size(),
        "qty_order": g["qty_order"].sum(),
        "qty_exec": g["qty_exec"].sum(),
        "qty_residual": g["qty_residual"].sum(),
        "qty_close_sent": g["qty_close_sent"].sum(),
        "qty_close_exec": g["qty_close_exec"].sum(),
        "unfilled": g["unfilled_reason"].apply(lambda s: s.ne("").sum()),
    }).to_string())
    log("")
    log("  side:  " + ", ".join(f"{k} {v:,}" for k, v in out["side"].value_counts().items()))
    log("")
    log("  how unfilled_reason was set:")
    log(audit["reason_rule"].value_counts().to_string())
    log("")

    log("  where the identifiers came from (orders, by market):")
    for col in ("sedol", "bbg_ticker", "ric"):
        src = audit[f"{col}_source"].replace("", "blank")
        log(f"    {col}:")
        log("      " + pd.crosstab(out["market"], src).to_string().replace("\n", "\n      "))
    log("")
    for col in ("sedol", "bbg_ticker", "ric"):
        blank = out[col].isna() | out[col].astype(str).str.strip().eq("")
        if blank.any():
            by = out.loc[blank, "market"].value_counts()
            warn(f"{col} blank on {int(blank.sum()):,} orders ("
                 + ", ".join(f"{k} {v}" for k, v in by.items()) + ")")
    for col in ("px_cont_last", "px_close", "volume_close"):
        n = int(out[col].isna().sum())
        if n:
            warn(f"{col} empty on {n:,} orders")
    over = int((out["qty_exec"] > out["qty_order"]).sum())
    if over:
        warn(f"qty_exec above qty_order on {over:,} orders; qty_residual set to 0 there")
    over = int((out["qty_close_exec"] > out["qty_close_sent"]).sum())
    if over:
        warn(f"qty_close_exec above qty_close_sent on {over:,} orders")
    over = int((out["qty_close_exec"] > out["volume_close"]).sum())
    if over:
        warn(f"qty_close_exec above volume_close on {over:,} orders - one of "
             f"the two is not auction-only")
    n = int(out["unfilled_reason"].eq(OTHER_UNEXPLAINED).sum())
    if n:
        warn(f"{n:,} orders carry '{OTHER_UNEXPLAINED}'. They are listed in "
             f"moc_close_audit.csv (reason_rule = unexplained). Replace before "
             f"sending.")
    log("")
    log("  Not computed, written empty: " + ", ".join(NOT_COMPUTED))
    log("  Assumptions to confirm before sending:")
    log("    - MOC orders are " + f"{sorted(MOC_ALGOS)}" + " on the algo column")
    log(f"    - qty_close_sent source: {CLOSE_SENT_SOURCE}")
    log("    - fillCloseSize and marketCloseSize are auction prints only (no TAL)")
    log("    - UPTICK_RULE: JP 10% price restriction, HK reference price ~ last_cont_price")
    if AUCTION_VOLUME_CAP_PCT is None:
        log("    - VOLUME_CAP never assigned: no participation cap configured")


def tca_section(title: str) -> None:
    log("")
    log("=" * 74)
    log(title)
    log("=" * 74)


def probe(aws: pd.DataFrame, cols: dict) -> None:
    tca_section("PROBE - nothing written except run_log.txt")
    log("  column mapping:")
    for k, cands in SOURCE.items():
        log(f"    {k:<14}{cols.get(k, '-- not found --'):<22}(looked for {', '.join(cands)})")
    for k, v in cols.items():
        if k in ("date", "sym", "side", "country", "end_time"):
            log(f"  sample {v}: {aws[v].dropna().astype(str).unique()[:6].tolist()}")

    market, disagree = market_of(aws, cols) if "sym" in cols else (pd.Series(dtype=object), 0)
    log("")
    log("  market (country, else symbol suffix):")
    log(market.fillna("unmapped").value_counts().to_string())
    if disagree:
        warn(f"{disagree:,} orders: country disagrees with symbol suffix")

    algo_col = find_column(aws.columns, ALGO_COLUMNS)
    if algo_col is not None and "fill_close" in cols:
        in_mkt = market.isin(scope_markets())
        sub = aws.loc[in_mkt]
        fill = num(sub, cols, "fill_close").fillna(0)
        cum = num(sub, cols, "cumqty")
        g = pd.DataFrame({"algo": sub[algo_col].astype(str),
                          "auction_share": fill / cum.where(cum > 0)}).groupby("algo")
        tab = pd.DataFrame({"orders": g.size(),
                            "mean auction share of executed": g["auction_share"].mean().round(3)})
        tab["in MOC_ALGOS"] = [a.strip().upper() in MOC_ALGOS for a in tab.index]
        log("")
        log(f"  {algo_col} values in {'/'.join(MARKETS)} (pin MOC_ALGOS from this):")
        log(tab.sort_values("orders", ascending=False).to_string())
    for extra in ("tag9001", "deskOptions", "auctionOnly"):
        c = find_column(aws.columns, [extra])
        if c is not None:
            log(f"  {c} top values: "
                f"{aws[c].astype(str).value_counts().head(8).to_dict()}")

    if "side" in cols:
        raw = aws[cols["side"]].astype(str).str.upper().str.replace(r"[\s_-]", "", regex=True)
        log("")
        log("  side codes -> requested value:")
        log(pd.DataFrame({"orders": raw.value_counts(),
                          "maps to": raw.value_counts().index.map(
                              lambda v: SIDE_MAP.get(v, "!! UNMAPPED"))}).to_string())

    if all(k in cols for k in ("ordqty", "cumqty", "fill_close")):
        o, c, f = (num(aws, cols, k) for k in ("ordqty", "cumqty", "fill_close"))
        resid = (o - (c - f.fillna(0))).clip(lower=0)
        log("")
        log("  qty_close_sent candidates (all in-window orders):")
        log(f"    residual = ordqty - (cumqty - fillCloseSize): median {resid.median():,.0f}")
        if "qatt_close" in cols:
            q = num(aws, cols, "qatt_close")
            both = q.notna() & resid.notna()
            log(f"    maxQattCloseSym: populated on {int(q.notna().sum()):,}, "
                f"median {q.median():,.0f}")
            if both.any():
                log(f"    equal on {int((np.isclose(q[both], resid[both])).sum()):,} "
                    f"of {int(both.sum()):,}; maxQatt below fillCloseSize on "
                    f"{int((q < f.fillna(0)).sum()):,}")


# ===========================================================================
# EXCEL
# ===========================================================================

TEXT_COLUMNS = {"trade_date", "sedol", "bbg_ticker", "ric", "aggrTgtId", "sym"}
QTY_COLUMNS = {"qty_order", "qty_exec", "qty_residual", "qty_close_sent",
               "qty_close_exec", "volume_close"}


def write_excel(out: pd.DataFrame, audit: pd.DataFrame, path: Path) -> None:
    """moc_close - the file to send, as the CSV - and audit, every order with
    its source columns and rules."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, frame in (("moc_close", out), ("audit", audit)):
            frame.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            if len(frame.columns):
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(frame.columns, start=1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDE4EE")
                cell.alignment = Alignment(vertical="center")
                fmt = ("@" if col in TEXT_COLUMNS else
                       "#,##0" if col in QTY_COLUMNS else None)
                if fmt:
                    for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        c.number_format = fmt
                width = max([len(str(col))] + [len(str(v)) for v in frame[col].head(500)])
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)


# ===========================================================================
# SELF-TEST
# ===========================================================================

def self_test() -> int:
    """Synthetic parquet across two dates, one row per rule."""
    base = dict(clientAlgo="CLOSE", country="Japan", ordqty=1000.0, cumqty=1000.0,
                fillCloseSize=1000.0, marketCloseSize=500000.0,
                last_cont_price=100.0, endprice=100.0, ordprice=0.0,
                PreviousClose=100.0, PX_LOW=99.0, fend_time="14:31:00",
                maxQattCloseSym=1000.0)
    rows = [
        # id, date, sym, side, overrides, expected reason, expected sent
        ("A1", "2026-07-21", "9984.T", "SSH",
         dict(cumqty=500, fillCloseSize=0, last_cont_price=5737, endprice=5751,
              ordprice=5752, PreviousClose=5700, PX_LOW=5600),
         "OTHER: limit price did not cross the auction price", 500),
        ("A2", "2026-07-21", "0700.HK", "Buy",
         dict(country="HK", cumqty=600, fillCloseSize=200, ordprice=100.0,
              fend_time="16:10:00"), "LIMIT_QUEUE", 600),
        ("A3", "2026-07-22", "BHP AT", "Sell",
         dict(country="Australia", fend_time="14:15:00"), "", 1000),
        ("A4", "2026-07-22", "7203 JT", "SSH",
         dict(cumqty=0, fillCloseSize=0, PreviousClose=110, PX_LOW=98,
              endprice=99, last_cont_price=99.5), "UPTICK_RULE", 1000),
        ("A5", "2026-07-22", "1299.HK", "SSH",
         dict(country="HK", cumqty=0, fillCloseSize=0, endprice=99,
              last_cont_price=99.5, fend_time="16:10:00"), "UPTICK_RULE", 1000),
        ("A6", "2026-07-23", "D05.SI", "Buy",
         dict(country="Singapore", cumqty=300, fillCloseSize=0,
              fend_time="15:00:00"),
         "OTHER: order ended before the closing auction", 0),
        ("A7", "2026-07-23", "6758.T", "BC",
         dict(cumqty=400, fillCloseSize=100), OTHER_UNEXPLAINED, 700),
        ("A8", "2026-07-23", "6758.T", "Buy",
         dict(cumqty=0, fillCloseSize=0, marketCloseSize=0),
         "OTHER: no closing auction held", 1000),
        ("A9", "2026-07-23", "OCBC.SP", "Buy",
         dict(country="Singapore", fend_time="17:10:00"), "", 1000),
        # out of scope: VWAP algo, another market, outside the window
        ("X1", "2026-07-23", "6758.T", "Buy", dict(clientAlgo="VWAP"), None, None),
        ("X2", "2026-07-23", "005930.KS", "Buy", dict(country="Korea"), None, None),
        ("X3", "2026-09-14", "6758.T", "Buy", {}, None, None),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp) / "aws"
        folder.mkdir()
        recs = []
        for oid, date, sym, side, over, _, _ in rows:
            r = dict(base, aggrTgtId=oid, _date=pd.Timestamp(date), sym=sym, side=side)
            r.update(over)
            recs.append(r)
        df = pd.DataFrame(recs)
        for i, (_, chunk) in enumerate(df.groupby(df["_date"].dt.date)):
            chunk.to_parquet(folder / f"part_{i}.parquet", index=False)
        # a duplicate order in a second file must not produce a second row
        df.iloc[[0]].to_parquet(folder / "dupe.parquet", index=False)

        aws = read_window(folder, pd.Timestamp(DATE_FROM), pd.Timestamp(DATE_TO))
        out, audit = build(aws, resolve(aws))

    fails = 0

    def check(name, got, want):
        nonlocal fails
        ok = got == want
        fails += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} {name}: got {got!r}" + ("" if ok else f", want {want!r}"))

    check("rows = in-scope orders", len(out), 9)
    check("columns in requested order", list(out.columns), OUTPUT_COLUMNS)
    by_id = audit.set_index("aggrTgtId")
    for oid, _, _, _, _, want_reason, want_sent in rows:
        if want_reason is None:
            check(f"{oid} excluded", oid in by_id.index, False)
            continue
        check(f"{oid} unfilled_reason", by_id.at[oid, "unfilled_reason"], want_reason)
        check(f"{oid} qty_close_sent", int(by_id.at[oid, "qty_close_sent"]), want_sent)
    check("A1 qty_exec is cumqty", int(by_id.at["A1", "qty_exec"]), 500)
    check("A1 qty_residual is qty_order - qty_exec", int(by_id.at["A1", "qty_residual"]), 500)
    check("A3 fully executed: qty_residual 0", int(by_id.at["A3", "qty_residual"]), 0)
    check("qty_exec and qty_residual follow qty_order",
          OUTPUT_COLUMNS[OUTPUT_COLUMNS.index("qty_order"):OUTPUT_COLUMNS.index("qty_order") + 3],
          ["qty_order", "qty_exec", "qty_residual"])
    check("A1 ric", by_id.at["A1", "ric"], "9984.T")
    check("A1 bbg", by_id.at["A1", "bbg_ticker"], "9984 JT")
    check("A1 side", by_id.at["A1", "side"], "SELL_SHORT")
    check("A2 bbg drops HK zeros", by_id.at["A2", "bbg_ticker"], "700 HK")
    check("A3 ric from bbg", by_id.at["A3", "ric"], "BHP.AX")
    check("A6 SG bbg not guessed", by_id.at["A6", "bbg_ticker"], None)
    check("A7 side", by_id.at["A7", "side"], "BUY_COVER")
    check("A2 date format", by_id.at["A2", "trade_date"], "2026-07-21")
    check("excluded columns empty", bool((out[NOT_COMPUTED] == "").all().all()), True)

    print("\nsym forms")
    check("house sym 9984.JP", identifiers("9984.JP"), ("9984.T", "9984 JT"))
    check("house sym 700.HK", identifiers("700.HK"), ("0700.HK", "700 HK"))
    check("house sym BHP.AU", identifiers("BHP.AU"), ("BHP.AX", "BHP AT"))
    check("house sym DBS.SP keeps bbg only", identifiers("DBS.SP"), (None, "DBS SP"))
    check("kdb candidates from a RIC", kdb_candidates("9984.T"), ["9984.T", "9984.JP"])
    check("kdb candidates for HK", kdb_candidates("0700.HK"), ["0700.HK", "700.HK"])

    print("\nkdb")

    class Result:
        def __init__(self, df):
            self.df = df

        def pd(self):
            return self.df

    print("\ncrosscode")
    cc_text = (
        "#FidessaCode,RicCode,Type,BloombergCode,FidessaMarket,BloombergStatus\n"
        # a delisted line first: the live one below must win
        "D05.SP,D05old.SI,Equity,DBS SP,SES-MAIN,DLST\n"
        "D05.SP,D05.SI,Equity,DBS SP,SES-MAIN,ACTV\n"
        "OCBC.SP,OCBC.SI,Equity,OCBC SP,SES-MAIN,ACTV\n"
        # a secondary Japanese venue, ignored
        "9984.JP,9984.JNX,Equity,9984 JT,JNX-MAIN,ACTV\n"
        "9984.JP,9984.T,Equity,9984 JT,TYO-MAIN,ACTV\n"
        "BHP.AU,BHP.AX,Equity,BHP AT,ASX-MAIN,ACTV\n"
        "LINK.HK,0823.HK,Equity,823 HK,HKG-MAIN,ACTV\n"
        "1299.HK,1299.HK,Equity,1299 HK,HKG-MAIN,\n"
        ",U11.SI,Equity,UOB SP,SES-MAIN,ACTV\n")
    with tempfile.TemporaryDirectory() as tmp:
        ccp = Path(tmp) / "CrossCode.csv"
        ccp.write_text(cc_text, encoding="utf-8")
        cc = load_crosscode(ccp)
        bad = Path(tmp) / "bad.csv"
        bad.write_text("FidessaCode,BloombergCode\nX.JP,X JT\n", encoding="utf-8")
        for name, path in (("a placeholder path stops the run", r"CHANGEME\CrossCode.csv"),
                           ("a CrossCode without RicCode stops the run", bad)):
            try:
                load_crosscode(path)
                check(name, False, True)
            except SystemExit:
                check(name, True, True)
    check("secondary venues are left out", "9984.JNX" in set(cc["RicCode"]), False)

    c_out, c_audit = out.copy(), audit.copy()
    apply_crosscode(c_out, c_audit, cc)
    cb = c_audit.set_index("aggrTgtId")
    check("A6 SG ric by RicCode", (cb.at["A6", "ric"], cb.at["A6", "cc_match"]),
          ("D05.SI", "RicCode"))
    check("A6 SG bbg from CrossCode",
          (cb.at["A6", "bbg_ticker"], cb.at["A6", "bbg_ticker_source"]),
          ("DBS SP", "crosscode"))
    check("A9 house sym OCBC.SP -> ric by FidessaCode",
          (cb.at["A9", "ric"], cb.at["A9", "cc_match"]), ("OCBC.SI", "FidessaCode"))
    check("A1 9984.T keeps the TYO line, not JNX", cb.at["A1", "ric"], "9984.T")
    check("A3 BHP AT by the RIC its sym converts to", cb.at["A3", "cc_match"],
          "RicCode from sym")
    t_out = pd.DataFrame({"market": ["SG"], "ric": [None], "bbg_ticker": ["UOB SP"]},
                         dtype=object)
    t_audit = pd.DataFrame({"sym": ["UOB SP"], "ric": [None], "bbg_ticker": ["UOB SP"],
                            "ric_source": [""], "bbg_ticker_source": ["extract"]})
    apply_crosscode(t_out, t_audit, cc)
    check("UOB SP, no FidessaCode or RIC to go on: by ticker within market",
          (t_audit.at[0, "cc_match"], t_out.at[0, "ric"]), ("ticker", "U11.SI"))
    check("A5 1299.HK, blank status kept", cb.at["A5", "ric_source"], "crosscode")
    check("A7 not in CrossCode keeps the sym ric",
          (cb.at["A7", "ric"], cb.at["A7", "ric_source"]), ("6758.T", "extract"))

    print("\nkdb")
    kdb_rows = pd.DataFrame([
        # trade-date row, found through the converted sym
        dict(date=pd.Timestamp("2026-07-21"), sym="9984.JP", ID_SEDOL1="6770620",
             TICKER="9984", EQY_PRIM_EXCH_SHRT="JT", ric_code="9984.T"),
        # only a day earlier than the order's trade date
        dict(date=pd.Timestamp("2026-07-21"), sym="BHP.AU", ID_SEDOL1="6144690",
             TICKER="BHP", EQY_PRIM_EXCH_SHRT="AT", ric_code="BHP.AX"),
        # SG: kdb keys DBS.SP, reachable from D05.SI only through CrossCode
        dict(date=pd.Timestamp("2026-07-23"), sym="DBS.SP", ID_SEDOL1="6175203",
             TICKER="DBS", EQY_PRIM_EXCH_SHRT="SP", ric_code="DBSM.SI"),
        # CrossCode's ric must survive a different kdb ric_code
        dict(date=pd.Timestamp("2026-07-23"), sym="OCBC.SP", ID_SEDOL1="6663689",
             TICKER="OCBC", EQY_PRIM_EXCH_SHRT="SP", ric_code="OCBC.XX"),
        dict(date=pd.Timestamp("2026-07-21"), sym="700.HK", ID_SEDOL1="BMMV2K8",
             TICKER="700", EQY_PRIM_EXCH_SHRT="HK", ric_code="0700.HK"),
    ])

    class FakeKdb:
        def __init__(self, fail_tables=()):
            self.calls, self.fail = [], fail_tables

        def __call__(self, q, a, b, s):
            self.calls.append((q, a, b, list(s)))
            if any(f" from {t} " in q for t in self.fail):
                raise RuntimeError("ric_code")
            keep = kdb_rows["sym"].isin(list(s))
            cols = ["date", "sym"] + [f for f in KDB_TABLES["equity"]]
            if " from equity_master " in q:
                cols.append("ric_code")
            return Result(kdb_rows.loc[keep, cols].reset_index(drop=True))

    syms = sorted({c for cands in row_candidates(c_audit) for c in cands})
    fk = FakeKdb()
    eq, table = fetch_equity(fk, pd.Timestamp(DATE_FROM), pd.Timestamp(DATE_TO), syms)
    k_out, k_audit = c_out.copy(), c_audit.copy()
    apply_kdb(k_out, k_audit, eq, table)
    kb = k_audit.set_index("aggrTgtId")
    check("equity_master answered first", table, "equity_master")
    check("dates cross as day counts from 2000.01.01",
          fk.calls[0][1:3], (int((pd.Timestamp(DATE_FROM) - Q_EPOCH).days),
                             int((pd.Timestamp(DATE_TO) - Q_EPOCH).days)))
    check("no `$ cast and no by in the query",
          ("`$" in fk.calls[0][0], " by " in fk.calls[0][0]), (False, False))
    check("A1 sedol from kdb", kb.at["A1", "sedol"], "6770620")
    check("A1 matched on the trade date", kb.at["A1", "kdb_match"], "trade date")
    check("A1 via the converted sym", kb.at["A1", "kdb_sym"], "9984.JP")
    check("A3 matched on another date", kb.at["A3", "kdb_match"], "other date")
    check("A3 sedol", kb.at["A3", "sedol"], "6144690")
    check("A6 SG reaches kdb through CrossCode",
          (kb.at["A6", "kdb_sym"], kb.at["A6", "sedol"]), ("DBS.SP", "6175203"))
    check("A6 CrossCode ric is not replaced by kdb", kb.at["A6", "ric"], "D05.SI")
    check("A9 CrossCode ric is not replaced by kdb", kb.at["A9", "ric"], "OCBC.SI")
    check("A2 no CrossCode line: kdb ric used",
          (kb.at["A2", "ric"], kb.at["A2", "ric_source"]), ("0700.HK", "kdb"))
    check("A2 HK through 700.HK", kb.at["A2", "sedol"], "BMMV2K8")
    check("A7 not in kdb keeps derived ric", kb.at["A7", "ric"], "6758.T")
    check("A7 flagged not found", kb.at["A7", "kdb_match"], "not found")
    check("out and audit agree", list(k_out["sedol"]), list(k_audit["sedol"]))

    fk = FakeKdb(fail_tables=("equity_master",))
    eq, table = fetch_equity(fk, pd.Timestamp(DATE_FROM), pd.Timestamp(DATE_TO), syms)
    check("equity_master failing falls back to equity", table, "equity")
    k_out, k_audit = c_out.copy(), c_audit.copy()
    apply_kdb(k_out, k_audit, eq, table)
    check("off equity, sedol still fills",
          k_audit.set_index("aggrTgtId").at["A1", "sedol"], "6770620")

    empty = FakeKdb(fail_tables=("equity_master", "equity"))
    try:
        fetch_equity(empty, pd.Timestamp(DATE_FROM), pd.Timestamp(DATE_TO), syms)
        check("no answer from either table stops the run", False, True)
    except SystemExit:
        check("no answer from either table stops the run", True, True)

    print("\nexcel")
    k_out.loc[k_out.index[0], "sedol"] = "0123456"
    with tempfile.TemporaryDirectory() as tmp:
        xp = Path(tmp) / "moc_close.xlsx"
        write_excel(k_out, k_audit, xp)
        book = pd.read_excel(xp, sheet_name=None, dtype={"sedol": str})
    check("two sheets, in order", list(book), ["moc_close", "audit"])
    check("moc_close has the requested columns", list(book["moc_close"].columns), OUTPUT_COLUMNS)
    check("moc_close has every order", len(book["moc_close"]), len(k_out))
    check("a sedol keeps its leading zero", book["moc_close"].at[0, "sedol"], "0123456")
    check("the workbook carries qty_exec and qty_residual",
          list(book["moc_close"]["qty_residual"])[:1] == [int(k_out["qty_residual"].iloc[0])],
          True)
    print(f"\nself-test: {'PASS' if not fails else f'{fails} FAILURE(S)'}")
    return 1 if fails else 0


# ===========================================================================

def main(argv=None) -> int:
    global CLOSE_SENT_SOURCE
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--out", type=Path, default=Path("output_moc_request"))
    p.add_argument("--from", dest="d_from", default=DATE_FROM)
    p.add_argument("--to", dest="d_to", default=DATE_TO)
    p.add_argument("--no-kdb", action="store_true",
                   help="skip kdb: sedol stays blank, tickers derived from sym")
    p.add_argument("--no-crosscode", action="store_true",
                   help="skip CrossCode: ric from kdb or the sym, SG may be blank")
    p.add_argument("--sent-source", choices=["residual", "qatt"],
                   default=CLOSE_SENT_SOURCE)
    p.add_argument("--probe", action="store_true",
                   help="print the mapping and the values to check, write nothing else")
    p.add_argument("--self-test", action="store_true")
    args = p.parse_args(argv)
    CLOSE_SENT_SOURCE = args.sent_source

    if args.self_test:
        return self_test()
    aws_dir = Path(AWS_DIR)
    if not aws_dir.is_dir():
        print(f"ERROR: AWS_DIR is not a folder: {aws_dir}", file=sys.stderr)
        return 2
    args.out.mkdir(parents=True, exist_ok=True)
    d_from, d_to = pd.Timestamp(args.d_from), pd.Timestamp(args.d_to)

    tca_section(f"MOC CLOSE REQUEST  {d_from:%Y-%m-%d} .. {d_to:%Y-%m-%d}  "
                f"{'/'.join(MARKETS)}")
    # Connect before reading the parquet, so a wrong server fails in seconds.
    conn = None if (args.no_kdb or args.probe) else kdb_connect()
    cc = None if (args.no_crosscode or args.probe) else load_crosscode(CROSSCODE_PATH)
    aws = read_window(aws_dir, d_from, d_to)
    try:
        if aws.empty:
            warn("no rows in the window")
            return 2
        cols = resolve(aws)
        if args.probe:
            probe(aws, cols)
            return 0

        tag = f"{d_from:%Y-%m-%d}_{d_to:%Y-%m-%d}"
        concat_path = args.out / f"aws_concat_{tag}.csv"
        aws.drop(columns=["trade_date_parsed"]).to_csv(concat_path, index=False)
        log(f"  wrote {concat_path}  ({len(aws):,} rows x {len(aws.columns) - 1} columns)")

        out, audit = build(aws, cols)
        tca_section("IDENTIFIERS")
        if cc is None:
            warn("--no-crosscode: ric comes from kdb or the sym only")
        else:
            apply_crosscode(out, audit, cc)
        if args.no_kdb:
            warn("--no-kdb: sedol is blank")
        else:
            syms = sorted({c for cands in row_candidates(audit) for c in cands})
            eq, table = fetch_equity(conn, d_from, d_to, syms)
            apply_kdb(out, audit, eq, table)
        report(out, audit)

        out_path = args.out / f"moc_close_{tag}.csv"
        out.to_csv(out_path, index=False)
        audit.to_csv(args.out / "moc_close_audit.csv", index=False)
        xlsx_path = args.out / f"moc_close_{tag}.xlsx"
        log("")
        log(f"  wrote {out_path}")
        log(f"  wrote {args.out / 'moc_close_audit.csv'}")
        log(f"  wrote {xlsx_path}")
        write_excel(out, audit, xlsx_path)
        return 0
    finally:
        (args.out / "run_log.txt").write_text("\n".join(_LOG), encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
